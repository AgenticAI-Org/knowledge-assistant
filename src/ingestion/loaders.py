"""Document loaders for PDF, DOCX, and TXT files under data/."""
from __future__ import annotations

import logging
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def load_pdf(path: Path, source: str) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": source, "page": page_num}))
    return docs


def load_docx(path: Path, source: str) -> list[Document]:
    from docx import Document as DocxDocument

    docx_file = DocxDocument(str(path))
    text = "\n".join(p.text for p in docx_file.paragraphs if p.text.strip())
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": source, "page": None})]


def load_txt(path: Path, source: str) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": source, "page": None})]


_LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".txt": load_txt}


def load_documents(data_dir: Path) -> tuple[list[Document], list[str]]:
    """Loads every supported file under data_dir, recursively.

    Returns (documents, failed_sources). Failures are returned rather than
    swallowed so the caller can refuse to replace a good index with a partial
    corpus -- see src/ingestion/indexer.py.

    `source` is the path relative to data_dir, so two same-named files in
    different subdirectories get distinct chunk IDs instead of colliding.
    """
    documents: list[Document] = []
    failed: list[str] = []
    if not data_dir.exists():
        raise FileNotFoundError(f"Data directory not found: {data_dir}")

    for path in sorted(data_dir.rglob("*")):
        if not path.is_file():
            continue
        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            logger.debug("Skipping unsupported file", extra={"file": path.name})
            continue
        source = path.relative_to(data_dir).as_posix()
        try:
            docs = _LOADERS[ext](path, source)
            if not docs:
                logger.warning("No extractable text found", extra={"file": source})
            documents.extend(docs)
        except Exception:
            logger.exception("Failed to load document", extra={"file": source})
            failed.append(source)

    return documents, failed
