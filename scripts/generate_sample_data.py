"""One-off generator for the sample policy corpus under data/.

Produces a small, deliberately multi-format (PDF + DOCX + TXT) set of
fictional company policy documents used to exercise the RAG pipeline:
hybrid search, memory follow-ups, citations, and at least one topic that
is intentionally NOT covered (sabbatical leave) to test the hallucination
guardrail.

Run: uv run python -m scripts.generate_sample_data
"""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from src.config import PROJECT_ROOT

DATA_DIR = PROJECT_ROOT / "data"


def write_pdf(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER)
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for heading, body in sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        for para in body.strip().split("\n\n"):
            story.append(Paragraph(para.strip(), styles["BodyText"]))
            story.append(Spacer(1, 6))
        story.append(Spacer(1, 10))
    doc.build(story)


def write_docx(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    doc = DocxDocument()
    doc.add_heading(title, level=0)
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        for para in body.strip().split("\n\n"):
            doc.add_paragraph(para.strip())
    doc.save(str(path))


def write_txt(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    lines = [title, "=" * len(title), ""]
    for heading, body in sections:
        lines.append(heading)
        lines.append("-" * len(heading))
        lines.append(body.strip())
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


LEAVE_POLICY = (
    "Leave Policy",
    [
        (
            "Annual Leave Entitlement",
            """
            Every full-time employee at Northwind Retail Corp. is entitled to 24 days of paid annual
            leave per calendar year, accrued at a rate of 2 days per completed month of service.
            Part-time employees accrue annual leave on a pro-rata basis according to their contracted
            weekly hours.

            Annual leave accrues from the date of joining, but newly hired employees may not take
            more than 5 days of annual leave during their probation period without written approval
            from their department head.
            """,
        ),
        (
            "Leave Carry-Forward",
            """
            Employees may carry forward a maximum of 8 unused annual leave days into the next
            calendar year. Carried-forward days must be used within the first quarter (January to
            March) of the new year, after which any remaining carried-forward balance is forfeited.
            Carry-forward requests above the 8-day cap require written approval from the department
            head and HR.

            Carried-forward days are always consumed before the current year's fresh entitlement,
            so that the oldest balance expires first. Employees can view their carry-forward balance
            and its expiry date on the HR portal leave dashboard.
            """,
        ),
        (
            "Sick Leave",
            """
            Employees are entitled to 12 days of paid sick leave per year. A medical certificate is
            required for sick leave exceeding 2 consecutive working days. Unused sick leave does not
            carry forward and is not encashed.

            Sick leave may also be used to care for an immediate family member, subject to manager
            approval and capped at 4 days per calendar year out of the 12-day entitlement.
            """,
        ),
        (
            "Parental Leave",
            """
            Birthing parents are entitled to 26 weeks of paid maternity leave, and non-birthing
            parents to 4 weeks of paid paternity leave, to be taken within 6 months of the child's
            birth or the date of adoption placement. Parental leave does not reduce the annual leave
            entitlement for that year.

            Employees intending to take parental leave should notify HR at least 8 weeks before the
            expected start date so that cover arrangements can be planned with the reporting manager.
            """,
        ),
        (
            "Bereavement and Emergency Leave",
            """
            Employees may take up to 5 days of paid bereavement leave on the death of an immediate
            family member (spouse, child, parent, or sibling), and up to 2 days for an extended
            family member. Bereavement leave is granted separately from annual and sick leave.
            """,
        ),
        (
            "Unpaid Leave",
            """
            Unpaid leave of up to 30 consecutive days may be granted at the company's discretion once
            an employee has exhausted their annual leave balance. Requests require the endorsement of
            the department head and final approval from HR. Health insurance coverage continues
            during approved unpaid leave, but annual leave does not accrue for that period.
            """,
        ),
        (
            "How to Apply for Leave",
            """
            Leave requests must be submitted through the HR portal at least 3 working days in advance
            for planned leave. Emergency leave should be reported to the employee's manager as soon as
            possible, with formal documentation submitted within 2 working days of returning to work.

            Managers are expected to approve or decline a leave request within 2 working days of
            submission. If a request is not actioned within 5 working days it is escalated
            automatically to the department head.
            """,
        ),
        (
            "Public Holidays",
            """
            Northwind Retail Corp. observes 12 public holidays per calendar year, published by HR
            each December for the year ahead. Employees required to work on a public holiday receive
            a compensatory day off, to be taken within 60 days of the holiday worked.
            """,
        ),
    ],
)

EMPLOYEE_HANDBOOK = (
    "Employee Handbook",
    [
        (
            "Code of Conduct",
            """
            All employees are expected to act with integrity, treat colleagues and customers with
            respect, and avoid conflicts of interest. Employees must not accept gifts worth more than
            INR 2,000 from vendors or clients without disclosing them to their manager. Harassment,
            discrimination, and retaliation of any kind will result in disciplinary action, up to and
            including termination.

            Any employee who becomes aware of a potential conflict of interest -- for example a family
            member employed by a supplier or competitor -- must declare it to HR in writing within 14
            days of becoming aware of it.
            """,
        ),
        (
            "Raising a Concern",
            """
            Concerns about misconduct can be raised with the reporting manager, with HR, or
            anonymously through the confidential ethics hotline. All reports are investigated by HR
            within 10 working days. The company prohibits retaliation against anyone who raises a
            concern in good faith, whether or not the concern is ultimately substantiated.
            """,
        ),
        (
            "Working Hours",
            """
            Standard working hours are 9:30 AM to 6:30 PM, Monday through Friday, with a one-hour
            lunch break. Northwind Retail Corp. operates a flexible working policy allowing employees
            to shift their start time between 8:00 AM and 10:30 AM, provided core hours (11:00 AM to
            4:00 PM) are respected.

            Employees are not expected to respond to work communications outside their working hours
            except when on a formally rostered on-call rotation.
            """,
        ),
        (
            "Benefits Overview",
            """
            Full-time employees are eligible for group health insurance covering the employee, spouse,
            and up to two dependent children, effective from the first day of employment. The company
            also provides a wellness allowance of INR 15,000 per year, reimbursable against gym
            memberships, fitness classes, or mental health counselling.

            Additional benefits include group term life insurance at three times annual base salary,
            an annual preventive health check-up, and an employee assistance programme offering
            confidential counselling sessions at no cost to the employee.
            """,
        ),
        (
            "Learning and Development",
            """
            Each employee has an annual learning budget of INR 40,000 that may be spent on courses,
            certifications, conferences, or books relevant to their role. Requests are approved by
            the reporting manager and reimbursed through the expense management system against
            itemized receipts. Certification exam fees are reimbursed only on a passing result.
            """,
        ),
        (
            "Probation and Confirmation",
            """
            New employees serve a probation period of 3 months from their date of joining. Performance
            is reviewed by the reporting manager at the end of the probation period, and confirmation
            of employment is communicated in writing by HR within 2 weeks of the review.

            Probation may be extended once, by a maximum of 3 additional months, where the reporting
            manager documents specific performance concerns and an improvement plan.
            """,
        ),
        (
            "Notice Period and Exit",
            """
            Confirmed employees are required to serve a notice period of 60 days on resignation;
            employees still on probation serve 15 days. The company may waive part of the notice
            period at its discretion. Exit formalities include return of all company IT assets, an
            exit interview with HR, and final settlement within 45 days of the last working day.
            """,
        ),
    ],
)

IT_POLICY = (
    "IT Policy",
    [
        (
            "Acceptable Use",
            """
            Company IT equipment and accounts must be used primarily for business purposes. Installing
            unauthorized software, disabling antivirus protection, or connecting personal storage
            devices to company laptops without IT approval is prohibited.

            Limited personal use of company devices is tolerated provided it does not interfere with
            work, consume significant bandwidth, or involve unlawful or offensive material.
            """,
        ),
        (
            "Password Standards",
            """
            Account passwords must be at least 12 characters long and combine upper and lower case
            letters, digits, and a symbol. Passwords expire every 90 days and the previous 5
            passwords cannot be reused. Sharing passwords with colleagues is prohibited under all
            circumstances, including with IT support staff, who will never ask for a password.
            """,
        ),
        (
            "VPN Access",
            """
            All remote access to internal systems must go through the corporate VPN. Employees can
            request VPN access via the IT Service Portal. If VPN access stops working, first confirm
            your internet connection is stable, then restart the VPN client.

            VPN sessions disconnect automatically after 12 hours or after 30 minutes of inactivity,
            whichever comes first, and must then be re-authenticated.
            """,
        ),
        (
            "Resetting Your VPN Password",
            """
            To reset a forgotten or expired VPN password, open the IT Service Portal, select
            'Password Reset', choose 'VPN Account', and verify your identity using your registered
            employee email and two-factor authentication code. The new password takes effect
            immediately and any active VPN sessions must be restarted. If the automated reset fails
            twice, raise a support ticket with the IT Helpdesk for a manual reset.
            """,
        ),
        (
            "Data Classification and Handling",
            """
            Company information is classified as Public, Internal, Confidential, or Restricted.
            Confidential and Restricted data must not be copied to personal devices, personal cloud
            storage, or personal email accounts. Restricted data, such as payroll records and
            customer payment information, may only be accessed from a company-managed device on the
            corporate VPN.
            """,
        ),
        (
            "Device Security and Loss",
            """
            Company laptops are issued with full-disk encryption enabled and must not be
            reconfigured to disable it. Screens must be locked when unattended. A lost or stolen
            device must be reported to the IT Helpdesk within 4 hours of discovery so the device can
            be remotely wiped and the associated credentials revoked.
            """,
        ),
        (
            "Reporting IT Issues",
            """
            Non-urgent IT issues (laptop hardware faults, software installation requests, access
            requests) should be logged as a support ticket through the IT Service Portal. Urgent
            issues affecting business operations (system outages, security incidents) should be
            reported immediately by calling the IT Helpdesk hotline.

            The IT Helpdesk hotline is staffed 24 hours a day. Non-urgent tickets are acknowledged
            within 1 working day and targeted for resolution within 3 working days.
            """,
        ),
    ],
)

TRAVEL_POLICY = (
    "Travel Policy",
    [
        (
            "Domestic Travel Booking",
            """
            All domestic business travel must be booked through the approved corporate travel portal
            at least 5 working days in advance where possible. Economy class is standard for flights
            under 4 hours; business class requires director-level approval.

            Travel booked outside the corporate portal is reimbursed only where the employee can
            demonstrate that the portal was unavailable or that the alternative booking was
            materially cheaper.
            """,
        ),
        (
            "International Travel",
            """
            International business travel requires approval from the department head and, for trips
            longer than 7 days, from the Finance director. Visa fees, travel insurance, and required
            vaccinations are paid by the company. Employees should apply for visas at least 3 weeks
            before departure.
            """,
        ),
        (
            "Accommodation Standards",
            """
            Hotel accommodation is capped at INR 6,000 per night in metro cities and INR 4,000 per
            night in non-metro cities, inclusive of taxes. Stays above the cap require prior written
            approval from the department head. Employees are expected to book refundable rates where
            the trip dates are not yet confirmed.
            """,
        ),
        (
            "Expense Reimbursement",
            """
            Travel expenses (meals, local transport, accommodation) must be submitted through the
            expense management system within 15 days of the trip's completion, accompanied by itemized
            receipts. Reimbursements are processed within 10 business days of approval.

            Expenses submitted more than 15 days after the trip require a written explanation and
            department head approval, and may be declined.
            """,
        ),
        (
            "Per Diem Rates",
            """
            Employees traveling domestically receive a per diem of INR 2,500 per day for meals and
            incidentals in metro cities, and INR 1,800 per day in non-metro cities. International per
            diem rates are published separately by the Finance team and vary by destination country.

            Per diem is paid for each full day of travel; partial days at the start and end of a trip
            are paid at half the applicable rate. No receipts are required for per diem itself.
            """,
        ),
        (
            "Travel Advances",
            """
            Employees may request a travel advance of up to 80 percent of estimated trip costs
            through the expense management system, no earlier than 7 days before departure. Any
            unused advance must be returned within 15 days of the trip's completion, at the same time
            as the expense claim is submitted.
            """,
        ),
    ],
)

REMOTE_WORK_POLICY = (
    "Remote Work Policy",
    [
        (
            "Eligibility",
            """
            Employees who have completed their probation period are eligible to work remotely for up
            to 10 working days per month, subject to reporting manager approval and the operational
            needs of the team. Roles requiring physical presence, such as warehouse and store
            operations, are not eligible for remote work.
            """,
        ),
        (
            "Requesting Remote Work",
            """
            Remote work days should be requested through the HR portal at least 2 working days in
            advance. Managers may require specific days of on-site presence for team meetings,
            planning sessions, or customer visits. Recurring remote work arrangements are reviewed
            every 6 months.
            """,
        ),
        (
            "Home Office Setup",
            """
            The company provides a one-time home office allowance of INR 20,000 for eligible remote
            workers, claimable against a desk, chair, monitor, or ergonomic accessories through the
            expense management system. The allowance may be claimed once every 3 years. Company
            laptops remain company property and must be returned on exit.
            """,
        ),
        (
            "Expectations While Remote",
            """
            Employees working remotely are expected to be reachable during core hours (11:00 AM to
            4:00 PM), to maintain a stable internet connection, and to join meetings with video
            enabled where practical. Remote work is not a substitute for childcare or dependent care
            arrangements.
            """,
        ),
        (
            "Working From Another City",
            """
            Employees who wish to work from a location other than their registered base city for more
            than 15 consecutive days must obtain written approval from HR, as this can carry tax and
            insurance implications. Working from outside the country requires prior clearance from
            both HR and the Finance team.
            """,
        ),
    ],
)

PERFORMANCE_POLICY = (
    "Performance and Compensation Policy",
    [
        (
            "Performance Review Cycle",
            """
            Northwind Retail Corp. runs two formal performance reviews each year: a mid-year check-in
            in July and a full annual review in January. Both reviews are conducted in the HR portal
            and combine employee self-assessment, reporting manager assessment, and peer feedback.
            """,
        ),
        (
            "Rating Scale",
            """
            Performance is rated on a five-point scale: Outstanding, Exceeds Expectations, Meets
            Expectations, Partially Meets Expectations, and Does Not Meet Expectations. Ratings are
            calibrated across each department before release to ensure consistency between managers.
            """,
        ),
        (
            "Annual Increments and Bonus",
            """
            Salary increments take effect from 1 April each year and are based on the January review
            rating, the employee's position in the salary band, and overall company performance. The
            annual performance bonus is paid with the April payroll and is prorated for employees who
            joined partway through the review year.
            """,
        ),
        (
            "Promotions",
            """
            Promotion decisions are made twice a year, alongside the January and July review cycles.
            Candidates are normally expected to have spent at least 18 months in their current role
            and to have received a rating of Exceeds Expectations or above in their most recent
            review. Promotions are recommended by the reporting manager and approved by the
            department head and HR.
            """,
        ),
        (
            "Performance Improvement Plans",
            """
            An employee rated Does Not Meet Expectations is placed on a formal Performance
            Improvement Plan of 60 days, with written objectives agreed between the employee, the
            reporting manager, and HR. Progress is reviewed fortnightly. Successful completion closes
            the plan; an unsuccessful outcome may lead to a role change or termination of employment.
            """,
        ),
        (
            "Payroll Schedule",
            """
            Salaries are credited on the last working day of each month. Payslips are available in
            the HR portal from the first working day of the following month. Payroll cut-off for
            changes such as bank details, tax declarations, or reimbursements is the 20th of the
            month.
            """,
        ),
    ],
)

FAQ_CONTENT = (
    "Company FAQs",
    [
        (
            "General",
            """
            Q: Who do I contact for payroll questions?
            A: Payroll queries should be directed to payroll@northwindretail.example or raised as a
            ticket in the HR portal under the 'Payroll' category.

            Q: How do I update my bank details?
            A: Bank detail updates can be submitted through the HR portal under 'My Profile > Payment
            Details'. Changes typically take effect from the next payroll cycle.

            Q: Where do I find my payslip?
            A: Payslips are published in the HR portal under 'My Profile > Payslips' from the first
            working day of each month, and can be downloaded as a PDF.

            Q: Who do I contact about my health insurance card?
            A: Insurance card queries go to the HR benefits desk at benefits@northwindretail.example.
            Replacement cards are typically issued within 10 working days.
            """,
        ),
        (
            "IT",
            """
            Q: How do I get a new laptop?
            A: Submit a hardware request ticket through the IT Service Portal. Standard laptop
            refresh cycle is every 3 years, or earlier in case of hardware failure verified by IT.

            Q: How do I request access to a shared drive?
            A: Raise an access request ticket in the IT Service Portal naming the drive and the
            business reason. Access requires approval from the data owner before IT can grant it.

            Q: My laptop is running slowly. What should I do?
            A: Restart the machine, confirm pending operating system updates are installed, and if
            the problem persists log a support ticket with the IT Helpdesk describing when it started.
            """,
        ),
        (
            "Workplace",
            """
            Q: How do I book a meeting room?
            A: Meeting rooms are booked through the calendar system by adding the room as a resource
            to the meeting invitation. Rooms held but unused for 15 minutes are released automatically.

            Q: What should I do if I lose my access badge?
            A: Report a lost badge to the facilities desk immediately so it can be deactivated. A
            replacement badge is issued within 2 working days.

            Q: Can I bring a visitor into the office?
            A: Visitors must be registered with the facilities desk at least 1 working day in advance
            and must be accompanied by their host at all times inside the building.
            """,
        ),
    ],
)


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    write_pdf(DATA_DIR / "Leave_Policy.pdf", *LEAVE_POLICY)
    write_pdf(DATA_DIR / "Employee_Handbook.pdf", *EMPLOYEE_HANDBOOK)
    write_pdf(DATA_DIR / "Performance_Policy.pdf", *PERFORMANCE_POLICY)
    write_docx(DATA_DIR / "IT_Policy.docx", *IT_POLICY)
    write_docx(DATA_DIR / "Travel_Policy.docx", *TRAVEL_POLICY)
    write_docx(DATA_DIR / "Remote_Work_Policy.docx", *REMOTE_WORK_POLICY)
    write_txt(DATA_DIR / "Company_FAQs.txt", *FAQ_CONTENT)

    print(f"Sample data written to {DATA_DIR}")
    for f in sorted(DATA_DIR.iterdir()):
        print(" -", f.name)


if __name__ == "__main__":
    main()
